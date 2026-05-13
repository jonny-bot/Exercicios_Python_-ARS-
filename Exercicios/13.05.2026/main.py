# ==========================================
# IMPORTS
# ==========================================

import tkinter as tk

from tkinter import (
    ttk,
    messagebox,
    filedialog
)

from openpyxl import Workbook
from openpyxl.styles import Font

# ==========================================
# CONFIGURAÇÕES
# ==========================================

ALIQUOTA_IR = 0.15

COR_FUNDO = "#0f172a"
COR_CARD = "#1e293b"
COR_TEXTO = "#f8fafc"
COR_DETALHE = "#38bdf8"
COR_BOTAO = "#0ea5e9"
COR_BOTAO_HOVER = "#0284c7"
COR_VERDE = "#22c55e"

# ==========================================
# FUNÇÕES
# ==========================================

def converter_taxa(taxa, tipo):

    if tipo == "Anual":

        return (
            ((1 + taxa / 100) ** (1 / 12)) - 1
        )

    return taxa / 100


def converter_periodo(periodo, tipo):

    return (
        periodo * 12
        if tipo == "Anos"
        else periodo
    )


def formatar_moeda(valor):

    return (
        f"R$ {valor:,.2f}"
        .replace(",", "X")
        .replace(".", ",")
        .replace("X", ".")
    )


def limpar_tabela():

    tabela.delete(
        *tabela.get_children()
    )


def validar_entradas(
    valor_inicial,
    aporte,
    taxa,
    periodo
):

    if valor_inicial < 0:
        raise ValueError

    if aporte < 0:
        raise ValueError

    if taxa < 0:
        raise ValueError

    if periodo <= 0:
        raise ValueError


# ==========================================
# EFEITOS BOTÃO
# ==========================================

def efeito_hover(event):

    botao.config(
        bg=COR_BOTAO_HOVER
    )


def sair_hover(event):

    botao.config(
        bg=COR_BOTAO
    )


# ==========================================
# EXPORTAR EXCEL
# ==========================================

def exportar_excel():

    if not tabela.get_children():

        messagebox.showwarning(
            "Aviso",
            "Nenhum dado para exportar."
        )

        return

    arquivo = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[
            (
                "Excel Files",
                "*.xlsx"
            )
        ],
        title="Salvar Planilha"
    )

    if not arquivo:
        return

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Simulação"

    # ==========================================
    # CABEÇALHOS
    # ==========================================

    for coluna_index, coluna in enumerate(
        colunas,
        start=1
    ):

        cell = sheet.cell(
            row=1,
            column=coluna_index
        )

        cell.value = coluna

        cell.font = Font(
            bold=True
        )

    # ==========================================
    # DADOS
    # ==========================================

    for row_index, item in enumerate(
        tabela.get_children(),
        start=2
    ):

        valores = tabela.item(item)["values"]

        for col_index, valor in enumerate(
            valores,
            start=1
        ):

            sheet.cell(
                row=row_index,
                column=col_index,
                value=valor
            )

    # ==========================================
    # AJUSTAR COLUNAS
    # ==========================================

    for coluna in sheet.columns:

        tamanho = 0

        letra_coluna = (
            coluna[0].column_letter
        )

        for cell in coluna:

            try:

                if (
                    len(str(cell.value))
                    > tamanho
                ):

                    tamanho = len(
                        str(cell.value)
                    )

            except:
                pass

        ajuste = tamanho + 5

        sheet.column_dimensions[
            letra_coluna
        ].width = ajuste

    workbook.save(arquivo)

    messagebox.showinfo(
        "Sucesso",
        "Planilha exportada com sucesso!"
    )


# ==========================================
# CALCULAR
# ==========================================

def calcular():

    try:

        valor_inicial = float(
            entry_valor_inicial.get()
        )

        aporte = float(
            entry_aporte.get()
        )

        taxa = float(
            entry_taxa.get()
        )

        periodo = int(
            entry_periodo.get()
        )

        validar_entradas(
            valor_inicial,
            aporte,
            taxa,
            periodo
        )

        meses = converter_periodo(
            periodo,
            combo_periodo.get()
        )

        taxa_mensal = converter_taxa(
            taxa,
            combo_taxa.get()
        )

        saldo = valor_inicial

        total_investido = valor_inicial

        juros_total = 0

        imposto_total = 0

        rendimento_liquido_total = 0

        limpar_tabela()

        for mes in range(1, meses + 1):

            saldo_anterior = saldo

            # ==========================================
            # BASE
            # ==========================================

            base_calculo = (
                saldo + aporte
            )

            # ==========================================
            # RENDIMENTO BRUTO
            # ==========================================

            rendimento_bruto = (
                base_calculo * taxa_mensal
            )

            # ==========================================
            # IMPOSTO
            # ==========================================

            imposto = (
                rendimento_bruto
                * ALIQUOTA_IR
            )

            # ==========================================
            # RENDIMENTO LÍQUIDO
            # ==========================================

            rendimento_liquido = (
                rendimento_bruto
                - imposto
            )

            # ==========================================
            # SALDO FINAL
            # ==========================================

            saldo = (
                base_calculo
                + rendimento_liquido
            )

            total_investido += aporte

            juros_total += rendimento_bruto

            imposto_total += imposto

            rendimento_liquido_total += (
                rendimento_liquido
            )

            # ==========================================
            # INSERIR TABELA
            # ==========================================

            tabela.insert(
                "",
                "end",
                values=(

                    mes,

                    formatar_moeda(
                        saldo_anterior
                    ),

                    formatar_moeda(
                        base_calculo
                    ),

                    formatar_moeda(
                        rendimento_bruto
                    ),

                    formatar_moeda(
                        imposto
                    ),

                    formatar_moeda(
                        rendimento_liquido
                    ),

                    formatar_moeda(
                        saldo
                    )
                )
            )

        atualizar_resultados(
            total_investido,
            juros_total,
            imposto_total,
            rendimento_liquido_total,
            saldo,
            taxa_mensal
        )

    except ValueError:

        messagebox.showerror(
            "Erro",
            "Digite valores válidos."
        )


# ==========================================
# RESULTADOS
# ==========================================

def atualizar_resultados(
    total_investido,
    juros_total,
    imposto_total,
    rendimento_liquido_total,
    saldo,
    taxa_mensal
):

    # ==========================================
    # RENDA MENSAL
    # ==========================================

    renda_mensal = (
        saldo * taxa_mensal
    ) * (1 - ALIQUOTA_IR)

    # ==========================================
    # LABELS
    # ==========================================

    label_investido.config(
        text=(
            f"Investido: "
            f"{formatar_moeda(total_investido)}"
        )
    )

    label_bruto.config(
        text=(
            f"Rendimento Bruto: "
            f"{formatar_moeda(juros_total)}"
        )
    )

    label_ir.config(
        text=(
            f"IR (15%): "
            f"{formatar_moeda(imposto_total)}"
        )
    )

    label_liquido.config(
        text=(
            f"Rendimento Líquido: "
            f"{formatar_moeda(rendimento_liquido_total)}"
        )
    )

    label_renda.config(
        text=(
            f"Renda Mensal Estimada: "
            f"{formatar_moeda(renda_mensal)}"
        )
    )

    label_total.config(
        text=(
            f"Saldo Final: "
            f"{formatar_moeda(saldo)}"
        )
    )


# ==========================================
# CRIAR CAMPOS
# ==========================================

def criar_campo(texto, coluna):

    label = tk.Label(
        frame_inputs,
        text=texto,
        font=(
            "Segoe UI",
            10,
            "bold"
        ),
        fg=COR_TEXTO,
        bg=COR_CARD
    )

    label.grid(
        row=0,
        column=coluna,
        sticky="w",
        padx=10,
        pady=(10, 5)
    )

    entrada = tk.Entry(
        frame_inputs,
        font=("Segoe UI", 11),
        relief="flat",
        bg="#334155",
        fg="white",
        insertbackground="white"
    )

    entrada.grid(
        row=1,
        column=coluna,
        sticky="ew",
        padx=10,
        pady=(0, 10),
        ipady=8
    )

    return entrada


# ==========================================
# JANELA
# ==========================================

janela = tk.Tk()

janela.title(
    "Simulador de Juros Compostos"
)

janela.state("zoomed")

janela.configure(
    bg=COR_FUNDO
)

# ==========================================
# GRID RESPONSIVO
# ==========================================

janela.grid_rowconfigure(
    4,
    weight=1
)

janela.grid_columnconfigure(
    0,
    weight=1
)

# ==========================================
# ESTILO
# ==========================================

style = ttk.Style()

style.theme_use("clam")

style.configure(
    "Treeview",
    background="white",
    foreground="black",
    fieldbackground="white",
    rowheight=32,
    font=("Segoe UI", 10)
)

style.configure(
    "Treeview.Heading",
    background="#dbeafe",
    foreground="black",
    font=("Segoe UI", 10, "bold")
)

# ==========================================
# HEADER
# ==========================================

frame_header = tk.Frame(
    janela,
    bg=COR_FUNDO
)

frame_header.grid(
    row=0,
    column=0,
    sticky="ew",
    padx=20,
    pady=(20, 10)
)

titulo = tk.Label(
    frame_header,
    text="Simulador de Juros Compostos",
    font=("Segoe UI", 28, "bold"),
    fg="white",
    bg=COR_FUNDO
)

titulo.pack()

subtitulo = tk.Label(
    frame_header,
    text="Planeje seus investimentos com projeção mensal",
    font=("Segoe UI", 11),
    fg="#94a3b8",
    bg=COR_FUNDO
)

subtitulo.pack()

# ==========================================
# INPUTS
# ==========================================

frame_inputs = tk.Frame(
    janela,
    bg=COR_CARD
)

frame_inputs.grid(
    row=1,
    column=0,
    sticky="ew",
    padx=20,
    pady=10
)

for i in range(6):

    frame_inputs.grid_columnconfigure(
        i,
        weight=1
    )

entry_valor_inicial = criar_campo(
    "Valor Inicial",
    0
)

entry_aporte = criar_campo(
    "Aporte Mensal",
    1
)

entry_taxa = criar_campo(
    "Taxa (%)",
    2
)

entry_periodo = criar_campo(
    "Tempo",
    4
)

# ==========================================
# COMBOBOX
# ==========================================

tk.Label(
    frame_inputs,
    text="Tipo Taxa",
    font=("Segoe UI", 10, "bold"),
    fg=COR_TEXTO,
    bg=COR_CARD
).grid(
    row=0,
    column=3,
    sticky="w",
    padx=10
)

combo_taxa = ttk.Combobox(
    frame_inputs,
    values=[
        "Mensal",
        "Anual"
    ],
    state="readonly"
)

combo_taxa.current(1)

combo_taxa.grid(
    row=1,
    column=3,
    sticky="ew",
    padx=10,
    ipady=5
)

tk.Label(
    frame_inputs,
    text="Período",
    font=("Segoe UI", 10, "bold"),
    fg=COR_TEXTO,
    bg=COR_CARD
).grid(
    row=0,
    column=5,
    sticky="w",
    padx=10
)

combo_periodo = ttk.Combobox(
    frame_inputs,
    values=[
        "Meses",
        "Anos"
    ],
    state="readonly"
)

combo_periodo.current(1)

combo_periodo.grid(
    row=1,
    column=5,
    sticky="ew",
    padx=10,
    ipady=5
)

# ==========================================
# BOTÕES
# ==========================================

frame_botao = tk.Frame(
    janela,
    bg=COR_FUNDO
)

frame_botao.grid(
    row=2,
    column=0,
    pady=15
)

botao = tk.Button(
    frame_botao,
    text="CALCULAR",
    command=calcular,
    bg=COR_BOTAO,
    fg="white",
    relief="flat",
    cursor="hand2",
    font=("Segoe UI", 12, "bold"),
    padx=40,
    pady=12
)

botao.pack(
    side="left",
    padx=10
)

botao.bind(
    "<Enter>",
    efeito_hover
)

botao.bind(
    "<Leave>",
    sair_hover
)

botao_excel = tk.Button(
    frame_botao,
    text="EXPORTAR EXCEL",
    command=exportar_excel,
    bg="#16a34a",
    fg="white",
    relief="flat",
    cursor="hand2",
    font=("Segoe UI", 11, "bold"),
    padx=30,
    pady=12
)

botao_excel.pack(
    side="left",
    padx=10
)

# ==========================================
# RESULTADOS
# ==========================================

frame_resultados = tk.Frame(
    janela,
    bg=COR_CARD
)

frame_resultados.grid(
    row=3,
    column=0,
    sticky="ew",
    padx=20,
    pady=10
)

label_investido = tk.Label(
    frame_resultados,
    text="Investido: R$ 0,00",
    fg="white",
    bg=COR_CARD,
    font=("Segoe UI", 12, "bold")
)

label_investido.pack(
    pady=5
)

label_bruto = tk.Label(
    frame_resultados,
    text="Rendimento Bruto: R$ 0,00",
    fg=COR_DETALHE,
    bg=COR_CARD,
    font=("Segoe UI", 12, "bold")
)

label_bruto.pack(
    pady=5
)

label_ir = tk.Label(
    frame_resultados,
    text="IR (15%): R$ 0,00",
    fg="#ef4444",
    bg=COR_CARD,
    font=("Segoe UI", 12, "bold")
)

label_ir.pack(
    pady=5
)

label_liquido = tk.Label(
    frame_resultados,
    text="Rendimento Líquido: R$ 0,00",
    fg="#facc15",
    bg=COR_CARD,
    font=("Segoe UI", 12, "bold")
)

label_liquido.pack(
    pady=5
)

label_renda = tk.Label(
    frame_resultados,
    text="Renda Mensal Estimada: R$ 0,00",
    fg="#38bdf8",
    bg=COR_CARD,
    font=("Segoe UI", 13, "bold")
)

label_renda.pack(
    pady=5
)

label_total = tk.Label(
    frame_resultados,
    text="Saldo Final: R$ 0,00",
    fg=COR_VERDE,
    bg=COR_CARD,
    font=("Segoe UI", 24, "bold")
)

label_total.pack(
    pady=10
)

# ==========================================
# TABELA
# ==========================================

frame_tabela = tk.Frame(
    janela,
    bg=COR_FUNDO
)

frame_tabela.grid(
    row=4,
    column=0,
    sticky="nsew",
    padx=20,
    pady=(10, 20)
)

frame_tabela.grid_rowconfigure(
    0,
    weight=1
)

frame_tabela.grid_columnconfigure(
    0,
    weight=1
)

colunas = (
    "Mês",
    "Saldo",
    "Base de Cálculo",
    "Rendimento Bruto",
    "IR (15%)",
    "Rendimento Líquido",
    "Saldo Final"
)

tabela = ttk.Treeview(
    frame_tabela,
    columns=colunas,
    show="headings",
    height=16
)

larguras = {
    "Mês": 80,
    "Saldo": 150,
    "Base de Cálculo": 170,
    "Rendimento Bruto": 180,
    "IR (15%)": 130,
    "Rendimento Líquido": 180,
    "Saldo Final": 180
}

for coluna in colunas:

    tabela.heading(
        coluna,
        text=coluna
    )

    tabela.column(
        coluna,
        width=larguras[coluna],
        anchor="center"
    )

scroll_y = ttk.Scrollbar(
    frame_tabela,
    orient="vertical",
    command=tabela.yview
)

scroll_x = ttk.Scrollbar(
    frame_tabela,
    orient="horizontal",
    command=tabela.xview
)

tabela.configure(
    yscrollcommand=scroll_y.set,
    xscrollcommand=scroll_x.set
)

tabela.grid(
    row=0,
    column=0,
    sticky="nsew"
)

scroll_y.grid(
    row=0,
    column=1,
    sticky="ns"
)

scroll_x.grid(
    row=1,
    column=0,
    sticky="ew"
)

# ==========================================
# EXECUTAR
# ==========================================

janela.mainloop()
