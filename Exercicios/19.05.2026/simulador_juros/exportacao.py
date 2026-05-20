import csv, json
from PyQt5.QtWidgets import QFileDialog, QMessageBox
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, numbers
from openpyxl.utils import get_column_letter
from datetime import datetime

def exportar_dados(widget):
    formatos = ["Excel (*.xlsx)", "CSV (*.csv)", "JSON (*.json)", "PDF (*.pdf)"]
    caminho, tipo = QFileDialog.getSaveFileName(widget, "Salvar arquivo", "", ";;".join(formatos))
    if not caminho:
        return

    dados = []
    for i in range(widget.tabela.rowCount()):
        linha = []
        for j in range(widget.tabela.columnCount()):
            item = widget.tabela.item(i, j)
            linha.append(item.text() if item else "")
        dados.append(linha)

    if tipo.startswith("Excel"):
        exportar_excel_formatado(widget, caminho, dados)

    elif tipo.startswith("CSV"):
        with open(caminho, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow([widget.tabela.horizontalHeaderItem(j).text() for j in range(widget.tabela.columnCount())])
            writer.writerows(dados)

    elif tipo.startswith("JSON"):
        colunas = [widget.tabela.horizontalHeaderItem(j).text() for j in range(widget.tabela.columnCount())]
        lista_dict = [dict(zip(colunas, linha)) for linha in dados]
        with open(caminho, "w", encoding="utf-8") as f:
            json.dump(lista_dict, f, ensure_ascii=False, indent=4)

    elif tipo.startswith("PDF"):
        # Usa ReportLab internamente; será lançada exceção se não instalado
        exportar_pdf_formatado(widget, caminho, [dict(zip([widget.tabela.horizontalHeaderItem(j).text() for j in range(widget.tabela.columnCount())], linha)) for linha in dados])

    QMessageBox.information(widget, "Exportação concluída", f"Arquivo salvo em:\n{caminho}")


def exportar_excel_formatado(widget, caminho, dados):
    """Exporta para Excel com formatação profissional"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Simulação"
    
    # Estilos
    titulo_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    titulo_font = Font(bold=True, color="FFFFFF", size=12)
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    
    borda = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Título
    end_col = get_column_letter(widget.tabela.columnCount())
    ws.merge_cells(f'A1:{end_col}1')
    titulo_cell = ws['A1']
    titulo_cell.value = widget.windowTitle()
    titulo_cell.font = titulo_font
    titulo_cell.fill = titulo_fill
    titulo_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 25
    
    # Data de exportação
    ws.merge_cells(f'A2:{end_col}2')
    data_cell = ws['A2']
    data_cell.value = f"Exportado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    data_cell.font = Font(italic=True, size=10)
    data_cell.alignment = Alignment(horizontal='center')
    
    # Cabeçalho
    colunas = [widget.tabela.horizontalHeaderItem(j).text() for j in range(widget.tabela.columnCount())]
    percent_columns = {idx for idx, titulo in enumerate(colunas, 1) if '%' in titulo or 'ir' in titulo.lower() or 'taxa' in titulo.lower()}
    for col_num, titulo_col in enumerate(colunas, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = titulo_col
        cell.font = header_font
        cell.fill = header_fill
        cell.border = borda
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    ws.row_dimensions[3].height = 20
    
    def parse_number(texto):
        texto = texto.strip().replace(" ", "")
        if texto == "":
            return None
        if "." in texto and "," in texto:
            if texto.rfind(".") > texto.rfind(","):
                texto = texto.replace(",", "")
            else:
                texto = texto.replace(".", "").replace(",", ".")
        elif "," in texto:
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
        try:
            return float(texto)
        except ValueError:
            return None

    def parse_valor(valor):
        texto = str(valor).strip()
        if texto == "":
            return texto, None
        is_percent = texto.endswith("%")
        is_currency = texto.startswith("R$")
        if is_currency:
            texto = texto.replace("R$", "").strip()
        if is_percent:
            texto = texto[:-1].strip()
        numero = parse_number(texto)
        if numero is None:
            return texto, None
        if is_percent:
            return numero / 100.0, 'percent'
        if is_currency:
            return numero, 'currency'
        return numero, 'number'

    # Dados
    for row_num, linha in enumerate(dados, 4):
        for col_num, valor in enumerate(linha, 1):
            cell = ws.cell(row=row_num, column=col_num)
            if col_num in percent_columns:
                texto = str(valor).strip()
                if texto == "":
                    cell.value = texto
                    tipo_valor = None
                else:
                    if texto.endswith("%"):
                        parsed = parse_number(texto[:-1])
                        cell.value = parsed / 100.0 if parsed is not None else texto
                        tipo_valor = 'percent' if parsed is not None else None
                    else:
                        parsed = parse_number(texto)
                        if parsed is not None:
                            cell.value = parsed / 100.0
                            tipo_valor = 'percent'
                        else:
                            cell.value, tipo_valor = parse_valor(valor)
            else:
                cell.value, tipo_valor = parse_valor(valor)

            cell.border = borda
            if col_num == 1:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='right', vertical='center')
            
            if tipo_valor == 'currency':
                cell.number_format = '_-"R$"* #,##0.00_-'
                cell.font = Font(color="000000")
            elif tipo_valor == 'percent':
                cell.number_format = '0.00%'
                cell.font = Font(color="000000")
            elif tipo_valor == 'number':
                cell.number_format = '#,##0.00'
    
    # Ajustar largura das colunas
    for col_num, titulo_col in enumerate(colunas, 1):
        col_letter = get_column_letter(col_num)
        if col_num == 1:
            ws.column_dimensions[col_letter].width = 12
        else:
            ws.column_dimensions[col_letter].width = 18
    
    # Adicionar sumário se houver atributos de totais
    sumario_row = len(dados) + 5
    
    if hasattr(widget, 'total_juros'):
        ws.row_dimensions[sumario_row].height = 20
        ws.merge_cells(f'A{sumario_row}:B{sumario_row}')
        label_cell = ws[f'A{sumario_row}']
        label_cell.value = "RESUMO FINAL"
        label_cell.font = Font(bold=True, size=11, color="FFFFFF")
        label_cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        label_cell.alignment = Alignment(horizontal='center', vertical='center')
        
        sumario_row += 1
        
        # Total em Juros
        ws[f'A{sumario_row}'].value = "Total em Juros:"
        ws[f'A{sumario_row}'].font = Font(bold=True)
        ws[f'B{sumario_row}'].value = widget.total_juros.text()
        if isinstance(ws[f'B{sumario_row}'].value, str) and ws[f'B{sumario_row}'].value.startswith('R$'):
            valor, tipo = parse_valor(ws[f'B{sumario_row}'].value)
            ws[f'B{sumario_row}'].value = valor
            ws[f'B{sumario_row}'].number_format = 'R$ #,##0.00'

        sumario_row += 1
        
        # Total Investido
        ws[f'A{sumario_row}'].value = "Total Investido:"
        ws[f'A{sumario_row}'].font = Font(bold=True)
        ws[f'B{sumario_row}'].value = widget.total_investido.text()
        if isinstance(ws[f'B{sumario_row}'].value, str) and ws[f'B{sumario_row}'].value.startswith('R$'):
            valor, tipo = parse_valor(ws[f'B{sumario_row}'].value)
            ws[f'B{sumario_row}'].value = valor
            ws[f'B{sumario_row}'].number_format = 'R$ #,##0.00'

        sumario_row += 1
        
        # Total Final
        ws[f'A{sumario_row}'].value = "Valor Final:"
        ws[f'A{sumario_row}'].font = Font(bold=True, color="FFFFFF")
        ws[f'A{sumario_row}'].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        ws[f'B{sumario_row}'].value = widget.total_final.text()
        if isinstance(ws[f'B{sumario_row}'].value, str) and ws[f'B{sumario_row}'].value.startswith('R$'):
            valor, tipo = parse_valor(ws[f'B{sumario_row}'].value)
            ws[f'B{sumario_row}'].value = valor
            ws[f'B{sumario_row}'].number_format = 'R$ #,##0.00'
        ws[f'B{sumario_row}'].font = Font(bold=True, color="FFFFFF")
        ws[f'B{sumario_row}'].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    
    wb.save(caminho)


def exportar_pdf_formatado(widget, caminho, dados):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    except Exception:
        raise RuntimeError('ReportLab não está instalado. Instale com: pip install reportlab')

    doc = SimpleDocTemplate(caminho, pagesize=A4, title="Relatório")
    elementos = []

    styles = getSampleStyleSheet()
    estilo_titulo = styles['Heading1']
    estilo_normal = styles['Normal']

    elementos.append(Paragraph('Relatório de Exportação', estilo_titulo))
    from datetime import datetime
    elementos.append(Paragraph(datetime.now().strftime('%d/%m/%Y %H:%M'), estilo_normal))
    elementos.append(Spacer(1, 12))

    if not dados:
        elementos.append(Paragraph('Sem dados para exportar.', estilo_normal))
        doc.build(elementos)
        return

    # Sumário estilizado com principais informações
    def parse_number_text(texto):
        if texto is None:
            return None
        s = str(texto).strip()
        if s == "":
            return None
        s = s.replace('R$', '').replace('%', '').strip()
        # normaliza separadores BR/EN
        if "." in s and "," in s:
            if s.rfind('.') > s.rfind(','):
                s = s.replace(',', '')
            else:
                s = s.replace('.', '').replace(',', '.')
        elif ',' in s:
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
        try:
            return float(s)
        except:
            return None

    # Campos principais vindos do widget quando disponíveis
    saldo_inicial = getattr(widget, 'valor_inicial', None)
    aporte_mensal = getattr(widget, 'aporte_mensal', None)
    taxa_juros = getattr(widget, 'taxa_juros', None)
    tipo_taxa = getattr(widget, 'tipo_taxa', None)
    tempo = getattr(widget, 'tempo', None)
    tipo_tempo = getattr(widget, 'tipo_tempo', None)

    resumo_items = []
    if saldo_inicial is not None:
        resumo_items.append(["Saldo Inicial:", saldo_inicial.text()])
    if aporte_mensal is not None:
        resumo_items.append(["Aporte Mensal:", aporte_mensal.text()])
    if taxa_juros is not None and tipo_taxa is not None:
        resumo_items.append(["Taxa de Juros:", f"{taxa_juros.text()} ({tipo_taxa.currentText()})"])
    if tempo is not None and tipo_tempo is not None:
        resumo_items.append(["Tempo:", f"{tempo.text()} {tipo_tempo.currentText()}"])

    # Valores finais dos cards
    try:
        total_investido_txt = widget.card_investido.valor_label.text()
        total_juros_txt = widget.card_juros.valor_label.text()
        total_final_txt = widget.card_final.valor_label.text()
    except Exception:
        total_investido_txt = total_juros_txt = total_final_txt = None

    if total_investido_txt:
        resumo_items.append(["Total Investido:", total_investido_txt])
    if total_juros_txt:
        resumo_items.append(["Total em Juros:", total_juros_txt])
    if total_final_txt:
        resumo_items.append(["Valor Final:", total_final_txt])

    if resumo_items:
        elementos.append(Paragraph('Resumo da Simulação', styles['Heading2']))
        # tabela de resumo compacta
        r_table = Table(resumo_items, colWidths=[200, 300])
        r_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F8FAFB')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('LEFTPADDING', (0,0), (-1,-1), 8),
            ('RIGHTPADDING', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        elementos.append(r_table)
        elementos.append(Spacer(1, 20))

    cabecalho = list(dados[0].keys())
    tabela = [cabecalho]
    for row in dados:
        linha = [row.get(h, '') for h in cabecalho]
        tabela.append(linha)

    t = Table(tabela, repeatRows=1)
    ts = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F81BD')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
        ('ALIGN', (0, 0), (0, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#CCCCCC')),
    ])

    for i in range(1, len(tabela)):
        if i % 2 == 0:
            ts.add('BACKGROUND', (0, i), (-1, i), colors.HexColor('#F2F2F2'))

    t.setStyle(ts)
    elementos.append(t)
    elementos.append(Spacer(1, 12))

    # Anexar gráficos ao final do PDF
    try:
        from graficos import grafico_evolucao, grafico_composicao
        import tempfile, os
    except Exception:
        grafico_evolucao = grafico_composicao = None

    if grafico_evolucao is not None:
        # Extrair séries a partir dos dados
        meses = []
        saldos = []
        rend_liq = []
        for row in dados:
            # detectar chaves mais prováveis
            keys = [k.lower() for k in cabecalho]
            # encontrar coluna mês
            mes_key = next((k for k in cabecalho if 'mês' in k.lower() or 'mes' in k.lower()), cabecalho[0])
            saldo_final_key = next((k for k in cabecalho if 'saldo final' in k.lower()), None)
            rendimento_liq_key = next((k for k in cabecalho if 'rendimento líq' in k.lower() or 'rendimento liq' in k.lower() or 'rendimento líquido' in k.lower() or 'rendimento' in k.lower()), None)

            meses.append(row.get(mes_key, ''))
            # saldo final preferencialmente
            if saldo_final_key and row.get(saldo_final_key):
                val = parse_number_text(row.get(saldo_final_key))
            else:
                # fallback: usar coluna 'Saldo' se existir
                saldo_key = next((k for k in cabecalho if k.lower().startswith('saldo')), None)
                val = parse_number_text(row.get(saldo_key)) if saldo_key else None
            saldos.append(val if val is not None else 0)

            if rendimento_liq_key and row.get(rendimento_liq_key):
                val2 = parse_number_text(row.get(rendimento_liq_key))
            else:
                val2 = parse_number_text(row.get(next((k for k in cabecalho if '%' in k or 'ir' in k.lower()), ''), ''))
            rend_liq.append(val2 if val2 is not None else 0)

        # calcular juros acumulados
        acumulado = 0.0
        juros_acum = []
        for v in rend_liq:
            acumulado += v
            juros_acum.append(acumulado)

        # gerar canvases e salvar temporariamente
        try:
            canvas1 = grafico_evolucao(meses, saldos, juros_acum)
            fig1 = canvas1.figure
            tmp1 = tempfile.NamedTemporaryFile(delete=False, suffix='.png')
            fig1.savefig(tmp1.name, dpi=150, bbox_inches='tight')
            tmp1.close()

            total_investido_val = parse_number_text(total_investido_txt) if total_investido_txt else 0
            total_juros_val = parse_number_text(total_juros_txt) if total_juros_txt else 0
            canvas2 = grafico_composicao(total_investido_val, total_juros_val)
            fig2 = canvas2.figure
            tmp2 = tempfile.NamedTemporaryFile(delete=False, suffix='.png')
            fig2.savefig(tmp2.name, dpi=150, bbox_inches='tight')
            tmp2.close()

            from reportlab.platypus import Image as RLImage
            elementos.append(Paragraph('Gráficos', styles['Heading2']))
            elementos.append(Spacer(1, 6))
            elementos.append(RLImage(tmp1.name, width=450, height=200))
            elementos.append(Spacer(1, 12))
            elementos.append(RLImage(tmp2.name, width=250, height=200))

            # remover arquivos temporários após geração do PDF (após build)
            def cleanup():
                try:
                    os.unlink(tmp1.name)
                    os.unlink(tmp2.name)
                except:
                    pass
        except Exception:
            pass

    doc.build(elementos)

    # Apagar temporários (se existirem)
    try:
        cleanup()
    except Exception:
        pass
